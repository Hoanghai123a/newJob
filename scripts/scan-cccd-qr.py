#!/usr/bin/env python3
"""
CCCD QR Scanner Migration Script

Phase 1: AUDIT - Tìm employment_histories thiếu thông tin, có cccd_version, thống kê theo factory
Phase 2: SCAN - Quét QR từ ảnh CCCD và xuất Excel để review trước khi import

Usage:
  python scan-cccd-qr.py --mode audit [--factory "Factory Name"]
  python scan-cccd-qr.py --mode scan-all --audit-file audit.xlsx
  python scan-cccd-qr.py --mode scan-file --input history-ids.txt
"""

import os
import sys
import argparse
import json
from pathlib import Path
from datetime import datetime
from typing import Optional, List, Dict, Any
from io import BytesIO

try:
    import requests
    from PIL import Image
    from pyzbar import pyzbar
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from dotenv import load_dotenv
    import cv2
    import numpy as np
except ImportError as e:
    print(f"Error: Missing required library: {e.name}")
    print("Install with: pip install -r scripts/requirements-qr.txt")
    sys.exit(1)

# Import deep scanner module
from cccd_deep_scanner import scan_cccd_qr_deep, ScanMode


# ============================================================================
# Configuration
# ============================================================================

def load_local_env():
    """Load .env file if it exists"""
    env_path = Path(".env")
    if env_path.exists():
        load_dotenv(env_path)


def get_config():
    """Get configuration from environment variables"""
    load_local_env()
    return {
        "pb_url": os.getenv("PB_URL") or os.getenv("VITE_PB_URL") or "http://127.0.0.1:8290",
        "pb_admin_email": os.getenv("PB_ADMIN_EMAIL"),
        "pb_admin_password": os.getenv("PB_ADMIN_PASSWORD"),
        "output_dir": Path(os.getenv("OUTPUT_DIR", "./output"))
    }


# ============================================================================
# PocketBase Connection
# ============================================================================

class PocketBaseClient:
    def __init__(self, base_url: str, email: str, password: str):
        self.base_url = base_url.rstrip("/")
        self.email = email
        self.password = password
        self.token = None

    def auth(self):
        """Authenticate as superuser"""
        url = f"{self.base_url}/api/collections/_superusers/auth-with-password"
        resp = requests.post(url, json={
            "identity": self.email,
            "password": self.password
        })
        resp.raise_for_status()
        data = resp.json()
        self.token = data["token"]
        return self.token

    def _headers(self):
        """Get authorization headers"""
        if not self.token:
            raise RuntimeError("Not authenticated. Call auth() first.")
        return {"Authorization": f"Bearer {self.token}"}

    def get_list(self, collection: str, **params):
        """Get paginated list of records"""
        url = f"{self.base_url}/api/collections/{collection}/records"
        resp = requests.get(url, headers=self._headers(), params=params)
        resp.raise_for_status()
        return resp.json()

    def get_full_list(self, collection: str, batch_size=500, **params):
        """Get all records with pagination"""
        all_items = []
        page = 1
        while True:
            result = self.get_list(collection, page=page, perPage=batch_size, **params)
            all_items.extend(result["items"])
            if page >= result["totalPages"]:
                break
            page += 1
        return all_items

    def get_file_url(self, record: Dict, filename: str) -> str:
        """Get URL for a file attachment"""
        collection_name = record.get("collectionName", "cccd_versions")
        record_id = record["id"]
        return f"{self.base_url}/api/files/{collection_name}/{record_id}/{filename}"

    def download_file(self, record: Dict, field_name: str) -> Optional[bytes]:
        """Download file from PocketBase record"""
        filename = record.get(field_name)
        if not filename:
            return None
        url = self.get_file_url(record, filename)
        resp = requests.get(url, headers=self._headers())
        if resp.status_code == 200:
            return resp.content
        return None


# ============================================================================
# Data Processing Utilities
# ============================================================================

def normalize_cccd(value: Any) -> str:
    """Remove all non-digit characters from CCCD"""
    return "".join(c for c in str(value or "") if c.isdigit())


def is_valid_cccd(value: Any) -> bool:
    """Check if CCCD is valid (9 or 12 digits)"""
    normalized = normalize_cccd(value)
    return len(normalized) in [9, 12]


def parse_date_ddmmyyyy(date_str: str) -> Optional[str]:
    """Parse DDMMYYYY format to YYYY-MM-DD"""
    if not date_str or len(date_str) != 8:
        return None
    try:
        day = date_str[0:2]
        month = date_str[2:4]
        year = date_str[4:8]
        return f"{year}-{month}-{day}"
    except:
        return None


def parse_qr_data(qr_text: str) -> Dict[str, str]:
    """Parse Vietnamese CCCD QR code format (pipe-delimited)"""
    parts = qr_text.split("|")
    if len(parts) < 7:
        return {"error": "Invalid QR format", "raw": qr_text}

    return {
        "cccd_number": normalize_cccd(parts[0].strip()),
        "old_cmnd": parts[1].strip() if len(parts) > 1 else "",
        "full_name": parts[2].strip() if len(parts) > 2 else "",
        "date_of_birth": parse_date_ddmmyyyy(parts[3].strip()) if len(parts) > 3 else "",
        "gender": parts[4].strip() if len(parts) > 4 else "",
        "address": parts[5].strip() if len(parts) > 5 else "",
        "issue_date": parse_date_ddmmyyyy(parts[6].strip()) if len(parts) > 6 else "",
        "expiry_date": parse_date_ddmmyyyy(parts[7].strip()) if len(parts) > 7 else "",
    }


# ============================================================================
# QR Code Scanner
# ============================================================================

def scan_qr_from_image(image_bytes: bytes, timeout_ms: int = 8000) -> Optional[Dict[str, str]]:
    """
    Scan QR code from image bytes using deep scanner
    Returns parsed CCCD data or None/error dict
    """
    try:
        # Use deep scanner with AUTO mode (balanced speed/accuracy)
        result = scan_cccd_qr_deep(
            image_bytes=image_bytes,
            mode=ScanMode.AUTO,
            timeout_ms=timeout_ms,
            progress_callback=lambda msg: print(f"    → {msg}")
        )

        if not result.success:
            if result.reason == "timeout":
                return {"error": "Scan timeout after 8s"}
            elif result.reason == "not_found":
                return None  # QR not found
            else:
                return {"error": result.reason}

        # Map deep scanner result to legacy format
        data = result.data
        return {
            "cccd_number": data.get("cccd", ""),
            "full_name": data.get("fullName", ""),
            "date_of_birth": data.get("dateOfBirth", ""),
            "gender": data.get("gender", ""),
            "address": data.get("address", ""),
            "issue_date": data.get("issuedDate", ""),
            "old_identity": data.get("oldIdentity", "")
        }

    except Exception as e:
        return {"error": str(e)}


# ============================================================================
# Phase 1: AUDIT
# ============================================================================

def run_audit_phase(pb: PocketBaseClient, factory_name: Optional[str], company_code: Optional[str], output_dir: Path):
    """
    Phase 1: Tìm employment_histories thiếu thông tin, có cccd_version VÀ có ảnh
    Thống kê theo factory và xuất file audit Excel
    """
    print("🔍 Phase 1: AUDIT - Đang quét employment_histories...")

    # Query employment_histories
    filter_parts = [
        "cccd_version != ''",  # Phải có cccd_version
        # Thiếu ít nhất 1 trong các field
        "(worker_cccd_snapshot = '' || worker_date_of_birth_snapshot = '' || cccd_issue_date = '' || worker_address_snapshot = '')"
    ]

    if company_code:
        # Query company by code first
        companies = pb.get_full_list("companies", filter=f'code = "{company_code}"')
        if not companies:
            print(f"❌ Không tìm thấy công ty với mã: {company_code}")
            sys.exit(1)
        company_id = companies[0]["id"]
        company_name = companies[0]["name"]
        print(f"✓ Tìm thấy công ty: {company_name} (ID: {company_id})")
        filter_parts.append(f'worker.tenant_company = "{company_id}"')

    if factory_name:
        filter_parts.append(f'factory.name = "{factory_name}"')

    filter_str = " && ".join(filter_parts)

    histories = pb.get_full_list(
        "employment_histories",
        filter=filter_str,
        expand="factory,worker,cccd_version",
        sort="created"
    )

    print(f"✓ Tìm thấy {len(histories)} lịch sử có cccd_version")

    # Lọc thêm: chỉ giữ những bản có ảnh trong cccd_version
    filtered_histories = []
    for hist in histories:
        cccd_version = hist.get("expand", {}).get("cccd_version")
        if cccd_version:
            has_image = cccd_version.get("front_image") or cccd_version.get("back_image")
            if has_image:
                filtered_histories.append(hist)

    histories = filtered_histories
    print(f"✓ Lọc còn {len(histories)} lịch sử có ảnh CCCD")

    # Thống kê theo factory
    factory_stats = {}
    details = []

    for hist in histories:
        factory = hist.get("expand", {}).get("factory", {})
        factory_id = hist.get("factory", "")
        factory_display = factory.get("name", "Unknown Factory")

        if factory_id not in factory_stats:
            factory_stats[factory_id] = {
                "name": factory_display,
                "count": 0,
                "missing_cccd": 0,
                "missing_dob": 0,
                "missing_issue_date": 0,
                "missing_address": 0
            }

        stats = factory_stats[factory_id]
        stats["count"] += 1

        # Đếm các field thiếu
        if not hist.get("worker_cccd_snapshot"):
            stats["missing_cccd"] += 1
        if not hist.get("worker_date_of_birth_snapshot"):
            stats["missing_dob"] += 1
        if not hist.get("cccd_issue_date"):
            stats["missing_issue_date"] += 1
        if not hist.get("worker_address_snapshot"):
            stats["missing_address"] += 1

        # Chi tiết từng record
        worker = hist.get("expand", {}).get("worker", {})
        details.append({
            "history_id": hist["id"],
            "uid": hist.get("uid", ""),
            "factory_name": factory_display,
            "worker_name": hist.get("worker_name_snapshot", ""),
            "worker_id": hist.get("worker", ""),
            "worker_uid": worker.get("uid", ""),
            "employee_code": hist.get("employee_code", ""),
            "join_date": hist.get("join_date", ""),
            "missing_cccd": "X" if not hist.get("worker_cccd_snapshot") else "",
            "missing_dob": "X" if not hist.get("worker_date_of_birth_snapshot") else "",
            "missing_issue_date": "X" if not hist.get("cccd_issue_date") else "",
            "missing_address": "X" if not hist.get("worker_address_snapshot") else "",
            "has_cccd_version": "✓" if hist.get("cccd_version") else "",
        })

    # Xuất Excel
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    filename = f"cccd-audit-{timestamp}.xlsx"
    output_path = output_dir / filename

    wb = Workbook()

    # Sheet 1: Summary
    ws_summary = wb.active
    ws_summary.title = "Summary"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    headers = ["Factory", "Total Records", "Missing CCCD", "Missing DOB", "Missing Issue Date", "Missing Address"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws_summary.cell(1, col_idx, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    row_idx = 2
    for factory_id, stats in factory_stats.items():
        ws_summary.cell(row_idx, 1, stats["name"])
        ws_summary.cell(row_idx, 2, stats["count"])
        ws_summary.cell(row_idx, 3, stats["missing_cccd"])
        ws_summary.cell(row_idx, 4, stats["missing_dob"])
        ws_summary.cell(row_idx, 5, stats["missing_issue_date"])
        ws_summary.cell(row_idx, 6, stats["missing_address"])
        row_idx += 1

    # Auto width
    for col_idx in range(1, len(headers) + 1):
        ws_summary.column_dimensions[get_column_letter(col_idx)].width = 20

    # Sheet 2: Details
    ws_details = wb.create_sheet("Details")

    detail_headers = [
        "History ID", "UID", "Factory", "Worker Name", "Worker ID", "Worker UID",
        "Employee Code", "Join Date", "Missing CCCD", "Missing DOB",
        "Missing Issue Date", "Missing Address", "Has CCCD Version"
    ]

    for col_idx, header in enumerate(detail_headers, 1):
        cell = ws_details.cell(1, col_idx, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, detail in enumerate(details, 2):
        ws_details.cell(row_idx, 1, detail["history_id"])
        ws_details.cell(row_idx, 2, detail["uid"])
        ws_details.cell(row_idx, 3, detail["factory_name"])
        ws_details.cell(row_idx, 4, detail["worker_name"])
        ws_details.cell(row_idx, 5, detail["worker_id"])
        ws_details.cell(row_idx, 6, detail["worker_uid"])
        ws_details.cell(row_idx, 7, detail["employee_code"])
        ws_details.cell(row_idx, 8, detail["join_date"])
        ws_details.cell(row_idx, 9, detail["missing_cccd"])
        ws_details.cell(row_idx, 10, detail["missing_dob"])
        ws_details.cell(row_idx, 11, detail["missing_issue_date"])
        ws_details.cell(row_idx, 12, detail["missing_address"])
        ws_details.cell(row_idx, 13, detail["has_cccd_version"])

    # Auto width
    for col_idx in range(1, len(detail_headers) + 1):
        ws_details.column_dimensions[get_column_letter(col_idx)].width = 15

    output_dir.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    print(f"\n✅ Audit hoàn tất!")
    print(f"📊 Tổng số: {len(histories)} records")
    print(f"🏭 Số factory: {len(factory_stats)}")
    print(f"📁 File lưu tại: {output_path}")

    return str(output_path)


# ============================================================================
# Phase 2: SCAN & FILL
# ============================================================================

def run_scan_phase(pb: PocketBaseClient, history_ids: List[str], output_dir: Path):
    """
    Phase 2: Quét QR từ ảnh CCCD, so sánh với DB, xuất Excel
    """
    print(f"🔍 Phase 2: SCAN - Đang xử lý {len(history_ids)} records...")

    results = []
    stats = {
        "total": len(history_ids),
        "scanned": 0,
        "qr_found": 0,
        "qr_failed": 0,
        "no_image": 0
    }

    for idx, history_id in enumerate(history_ids, 1):
        print(f"  [{idx}/{len(history_ids)}] Processing {history_id}...")

        try:
            # Query history with expansions
            hist_result = pb.get_list(
                "employment_histories",
                filter=f'id="{history_id}"',
                expand="factory,worker,cccd_version"
            )

            if not hist_result["items"]:
                results.append({
                    "history_id": history_id,
                    "error": "History not found"
                })
                continue

            hist = hist_result["items"][0]
            factory = hist.get("expand", {}).get("factory", {})
            worker = hist.get("expand", {}).get("worker", {})
            cccd_version = hist.get("expand", {}).get("cccd_version")

            result = {
                "history_id": hist["id"],
                "uid": hist.get("uid", ""),
                "factory_name": factory.get("name", ""),
                "worker_name": hist.get("worker_name_snapshot", ""),
                "worker_uid": worker.get("uid", ""),
                "employee_code": hist.get("employee_code", ""),
                "join_date": hist.get("join_date", ""),

                # Database values
                "cccd_db": hist.get("worker_cccd_snapshot", ""),
                "dob_db": hist.get("worker_date_of_birth_snapshot", ""),
                "address_db": hist.get("worker_address_snapshot", ""),
                "issue_date_db": hist.get("cccd_issue_date", ""),

                # QR values (will be filled)
                "cccd_qr": "",
                "name_qr": "",
                "dob_qr": "",
                "gender_qr": "",
                "address_qr": "",
                "issue_date_qr": "",

                # Comparison
                "cccd_match": "",
                "dob_match": "",
                "address_match": "",
                "issue_date_match": "",

                # Status
                "qr_status": "",
                "action": "",
                "errors": ""
            }

            if not cccd_version:
                result["errors"] = "No CCCD version"
                result["qr_status"] = "NO_VERSION"
                stats["no_image"] += 1
                results.append(result)
                continue

            # Try front_image first, then back_image if QR not found
            qr_data = None
            image_side = None

            # Try front image first
            if cccd_version.get("front_image"):
                print(f"  → Scanning front image...")
                front_data = pb.download_file(cccd_version, "front_image")
                if front_data:
                    qr_data = scan_qr_from_image(front_data, timeout_ms=8000)
                    stats["scanned"] += 1

                    # Check if QR found successfully
                    if qr_data and not qr_data.get("error"):
                        image_side = "front"
                        print(f"  ✓ QR found on front image")

            # If no QR found on front, try back image
            if not qr_data or qr_data.get("error"):
                if cccd_version.get("back_image"):
                    print(f"  → Scanning back image...")
                    back_data = pb.download_file(cccd_version, "back_image")
                    if back_data:
                        qr_data = scan_qr_from_image(back_data, timeout_ms=8000)
                        stats["scanned"] += 1

                        if qr_data and not qr_data.get("error"):
                            image_side = "back"
                            print(f"  ✓ QR found on back image")

            # Check if we have any images at all
            if not cccd_version.get("front_image") and not cccd_version.get("back_image"):
                result["errors"] = "No image available"
                result["qr_status"] = "NO_IMAGE"
                stats["no_image"] += 1
                results.append(result)
                continue

            if not qr_data or "error" in qr_data:
                error_msg = qr_data.get("error", "QR not found") if qr_data else "QR not found"
                result["errors"] = f"Scan failed ({image_side}): {error_msg}"
                result["qr_status"] = "SCAN_FAILED"
                stats["qr_failed"] += 1
                results.append(result)
                continue

            # Fill QR data
            result["cccd_qr"] = qr_data.get("cccd_number", "")
            result["name_qr"] = qr_data.get("full_name", "")
            result["dob_qr"] = qr_data.get("date_of_birth", "")
            result["gender_qr"] = qr_data.get("gender", "")
            result["address_qr"] = qr_data.get("address", "")
            result["issue_date_qr"] = qr_data.get("issue_date", "")
            result["qr_status"] = f"SUCCESS ({image_side})"
            stats["qr_found"] += 1

            # Compare
            result["cccd_match"] = "✓" if normalize_cccd(result["cccd_db"]) == normalize_cccd(result["cccd_qr"]) else "✗"
            result["dob_match"] = "✓" if result["dob_db"] == result["dob_qr"] else "✗"
            result["address_match"] = "✓" if result["address_db"] == result["address_qr"] else "✗"
            result["issue_date_match"] = "✓" if result["issue_date_db"] == result["issue_date_qr"] else "✗"

            # Determine action
            needs_update = (
                not result["cccd_db"] or result["cccd_match"] == "✗" or
                not result["dob_db"] or result["dob_match"] == "✗" or
                not result["address_db"] or result["address_match"] == "✗" or
                not result["issue_date_db"] or result["issue_date_match"] == "✗"
            )
            result["action"] = "UPDATE" if needs_update else "SKIP"

            results.append(result)

        except Exception as e:
            results.append({
                "history_id": history_id,
                "error": str(e),
                "qr_status": "ERROR"
            })
            stats["qr_failed"] += 1

    # Export to Excel
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    filename = f"cccd-scan-result-{timestamp}.xlsx"
    output_path = output_dir / filename

    wb = Workbook()

    # ========== Sheet 1: Summary Statistics ==========
    ws_summary = wb.active
    ws_summary.title = "Summary"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    red_fill_light = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    # Summary headers
    ws_summary.cell(1, 1, "Thống kê kết quả quét QR").font = Font(bold=True, size=14)

    row = 3
    ws_summary.cell(row, 1, "Tổng số records").font = Font(bold=True)
    ws_summary.cell(row, 2, stats['total'])

    row += 1
    ws_summary.cell(row, 1, "✅ Quét được QR và điền được thông tin").font = Font(bold=True)
    ws_summary.cell(row, 2, stats['qr_found']).fill = green_fill

    row += 1
    ws_summary.cell(row, 1, "❌ Không quét được QR").font = Font(bold=True)
    ws_summary.cell(row, 2, stats['qr_failed']).fill = red_fill_light

    row += 1
    ws_summary.cell(row, 1, "⚠️  Không có ảnh CCCD").font = Font(bold=True)
    ws_summary.cell(row, 2, stats['no_image']).fill = yellow_fill

    row += 2
    ws_summary.cell(row, 1, "Tỷ lệ thành công").font = Font(bold=True)
    success_rate = f"{(stats['qr_found'] / stats['total'] * 100) if stats['total'] > 0 else 0:.1f}%"
    ws_summary.cell(row, 2, success_rate).fill = green_fill

    # Breakdown by status
    row += 2
    ws_summary.cell(row, 1, "Chi tiết trạng thái:").font = Font(bold=True, size=12)

    status_counts = {}
    for result in results:
        status = result.get("qr_status", "UNKNOWN")
        status_counts[status] = status_counts.get(status, 0) + 1

    row += 1
    ws_summary.cell(row, 1, "Trạng thái").font = header_font
    ws_summary.cell(row, 1).fill = header_fill
    ws_summary.cell(row, 2, "Số lượng").font = header_font
    ws_summary.cell(row, 2).fill = header_fill

    for status, count in sorted(status_counts.items()):
        row += 1
        ws_summary.cell(row, 1, status)
        ws_summary.cell(row, 2, count)
        if "SUCCESS" in status:
            ws_summary.cell(row, 2).fill = green_fill
        elif "FAILED" in status or "ERROR" in status:
            ws_summary.cell(row, 2).fill = red_fill_light
        else:
            ws_summary.cell(row, 2).fill = yellow_fill

    ws_summary.column_dimensions['A'].width = 40
    ws_summary.column_dimensions['B'].width = 15

    # ========== Sheet 2: Scan Results ==========
    ws = wb.create_sheet("Scan Results")

    headers = [
        "History ID", "UID", "Factory", "Worker Name", "Worker UID", "Employee Code", "Join Date",
        "CCCD DB", "CCCD QR", "CCCD Match",
        "DOB DB", "DOB QR", "DOB Match",
        "Address DB", "Address QR", "Address Match",
        "Issue Date DB", "Issue Date QR", "Issue Date Match",
        "Name QR", "Gender QR",
        "QR Status", "Action", "Errors"
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    for row_idx, result in enumerate(results, 2):
        ws.cell(row_idx, 1, result.get("history_id", ""))
        ws.cell(row_idx, 2, result.get("uid", ""))
        ws.cell(row_idx, 3, result.get("factory_name", ""))
        ws.cell(row_idx, 4, result.get("worker_name", ""))
        ws.cell(row_idx, 5, result.get("worker_uid", ""))
        ws.cell(row_idx, 6, result.get("employee_code", ""))
        ws.cell(row_idx, 7, result.get("join_date", ""))

        ws.cell(row_idx, 8, result.get("cccd_db", ""))
        ws.cell(row_idx, 9, result.get("cccd_qr", ""))
        match_cell = ws.cell(row_idx, 10, result.get("cccd_match", ""))
        if result.get("cccd_match") == "✗":
            match_cell.fill = red_fill

        ws.cell(row_idx, 11, result.get("dob_db", ""))
        ws.cell(row_idx, 12, result.get("dob_qr", ""))
        match_cell = ws.cell(row_idx, 13, result.get("dob_match", ""))
        if result.get("dob_match") == "✗":
            match_cell.fill = red_fill

        ws.cell(row_idx, 14, result.get("address_db", ""))
        ws.cell(row_idx, 15, result.get("address_qr", ""))
        match_cell = ws.cell(row_idx, 16, result.get("address_match", ""))
        if result.get("address_match") == "✗":
            match_cell.fill = red_fill

        ws.cell(row_idx, 17, result.get("issue_date_db", ""))
        ws.cell(row_idx, 18, result.get("issue_date_qr", ""))
        match_cell = ws.cell(row_idx, 19, result.get("issue_date_match", ""))
        if result.get("issue_date_match") == "✗":
            match_cell.fill = red_fill

        ws.cell(row_idx, 20, result.get("name_qr", ""))
        ws.cell(row_idx, 21, result.get("gender_qr", ""))

        ws.cell(row_idx, 22, result.get("qr_status", ""))
        ws.cell(row_idx, 23, result.get("action", ""))
        error_cell = ws.cell(row_idx, 24, result.get("errors", ""))
        if result.get("errors"):
            error_cell.fill = red_fill

    # Auto width
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 15

    output_dir.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)

    print(f"\n✅ Scan hoàn tất!")
    print(f"📊 Thống kê:")
    print(f"  - Total: {stats['total']}")
    print(f"  - Scanned: {stats['scanned']}")
    print(f"  - QR found: {stats['qr_found']}")
    print(f"  - QR failed: {stats['qr_failed']}")
    print(f"  - No image: {stats['no_image']}")
    print(f"📁 File lưu tại: {output_path}")

    return str(output_path)


# ============================================================================
# Main Entry Point
# ============================================================================

def parse_args():
    parser = argparse.ArgumentParser(description="CCCD QR Scanner Migration Script")
    parser.add_argument("--mode", required=True, choices=["audit", "scan-all", "scan-file"],
                       help="Mode: audit | scan-all | scan-file")
    parser.add_argument("--company", help="Company code (e.g., HL, HRP) to filter by tenant_company")
    parser.add_argument("--factory", help="Factory name (for audit mode)")
    parser.add_argument("--audit-file", help="Audit Excel file (for scan-all mode)")
    parser.add_argument("--input", help="Input file with history IDs (for scan-file mode)")
    return parser.parse_args()


def main():
    args = parse_args()
    config = get_config()

    # Validate config
    if not config["pb_admin_email"] or not config["pb_admin_password"]:
        print("Error: Missing PB_ADMIN_EMAIL or PB_ADMIN_PASSWORD in environment/.env")
        sys.exit(1)

    # Connect to PocketBase
    print(f"🔗 Connecting to PocketBase at {config['pb_url']}...")
    pb = PocketBaseClient(config["pb_url"], config["pb_admin_email"], config["pb_admin_password"])
    pb.auth()
    print("✓ Authenticated")

    # Run appropriate phase
    if args.mode == "audit":
        run_audit_phase(pb, args.factory, args.company, config["output_dir"])

    elif args.mode == "scan-all":
        if not args.audit_file:
            print("Error: --audit-file required for scan-all mode")
            sys.exit(1)

        # Resolve wildcard if present
        audit_file_path = Path(args.audit_file)
        if "*" in args.audit_file:
            import glob
            matches = glob.glob(args.audit_file)
            if not matches:
                print(f"Error: No files match pattern: {args.audit_file}")
                sys.exit(1)
            if len(matches) > 1:
                print(f"Error: Multiple files match pattern: {args.audit_file}")
                print("Matches:")
                for m in matches:
                    print(f"  - {m}")
                print("Please specify exact file path.")
                sys.exit(1)
            audit_file_path = Path(matches[0])
            print(f"📂 Resolved to: {audit_file_path}")

        if not audit_file_path.exists():
            print(f"Error: File not found: {audit_file_path}")
            sys.exit(1)

        # Read history IDs from audit file (Details sheet, column A)
        from openpyxl import load_workbook
        wb = load_workbook(str(audit_file_path))
        ws = wb["Details"]
        history_ids = [ws.cell(row, 1).value for row in range(2, ws.max_row + 1) if ws.cell(row, 1).value]
        wb.close()

        print(f"📋 Loaded {len(history_ids)} history IDs from {audit_file_path}")
        run_scan_phase(pb, history_ids, config["output_dir"])

    elif args.mode == "scan-file":
        if not args.input:
            print("Error: --input required for scan-file mode")
            sys.exit(1)

        # Read history IDs from text file (one per line)
        input_path = Path(args.input)
        if not input_path.exists():
            print(f"Error: File not found: {args.input}")
            sys.exit(1)

        history_ids = [line.strip() for line in input_path.read_text().splitlines() if line.strip()]
        print(f"📋 Loaded {len(history_ids)} history IDs from {args.input}")
        run_scan_phase(pb, history_ids, config["output_dir"])


if __name__ == "__main__":
    main()
